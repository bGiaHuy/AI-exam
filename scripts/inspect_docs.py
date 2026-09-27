import sys
import io
import os
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def analyze_docx(path):
    print(f"==================================================")
    print(f"FILE: {path}")
    print(f"==================================================")
    if not os.path.exists(path):
        print("File does not exist!")
        return
    doc = docx.Document(path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    print(f"Total sections: {len(doc.sections)}")
    
    # Check styles
    styles = set(p.style.name for p in doc.paragraphs)
    print(f"Paragraph styles used: {styles}")
    
    # List paragraphs
    print("\n--- OUTLINE / STRUCTURE ---")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            # Check if header or section
            if p.style.name.startswith('Heading') or any(text.startswith(prefix) for prefix in [
                'PHẦN', 'Phần', 'MỤC LỤC', 'DANH MỤC', 'TÊN ĐỀ TÀI', 'LỜI CẢM ƠN', 
                'CHƯƠNG', 'I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', '1.', '2.', '3.', '4.', '5.'
            ]):
                print(f"[{p.style.name}] (p.{i}): {text[:100]}")
    
    # Tables summary
    print(f"\n--- TABLES ({len(doc.tables)}) ---")
    for t_idx, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns) if rows > 0 else 0
        first_row_text = [cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells] if rows > 0 else []
        print(f"Table {t_idx+1}: {rows} rows x {cols} cols | Header: {first_row_text[:5]}")

if __name__ == '__main__':
    files = [
        'docx/BAO_CAO_KHKT_CAP_TINH_HE_THONG_GIAM_SAT_PHONG_THI_V0_4_1790334501.docx',
        'docx/Gian lan thi cu.docx',
        'Bản sao của BÁO CÁO KHKT.docx'
    ]
    for f in files:
        analyze_docx(f)
