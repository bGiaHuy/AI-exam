from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = Path(__file__).resolve().parent
STYLE_TOOLS = ROOT / "tools" / "vietnamese-docs-style" / "scripts"
sys.path.insert(0, str(STYLE_TOOLS))
from build_docx import setup_document  # noqa: E402


OUTPUT = OUT_DIR / "AI_EXAM_CONTROL_GIAO_AN_CHUYEN_NGANH_5_BUOI.docx"
BLACK = RGBColor(0, 0, 0)
GRAY = "E7E6E6"

SECTIONS = [
    ("PHẦN I. ĐỊNH HƯỚNG VÀ BẢN ĐỒ KIẾN THỨC", ["README.md", "02_KHUNG_KIEN_THUC.md"]),
    ("PHẦN II. GIÁO ÁN CHI TIẾT 5 BUỔI", ["01_GIAO_AN.md"]),
    ("PHẦN III. SỔ TAY HỌC VIÊN", ["03_SO_TAY_HOC_VIEN.md"]),
    ("PHẦN IV. THỰC HÀNH VÀ ĐÁNH GIÁ", ["12_PHIEU_THUC_HANH.md", "07_KIEM_TRA.md", "10_PHIEU_DANH_GIA.md"]),
    ("PHẦN V. TRÌNH DIỄN, PHẢN BIỆN VÀ VẬN HÀNH", ["04_KICH_BAN_DEMO.md", "05_NGAN_HANG_CAU_HOI.md", "06_DAP_AN_MAU.md", "08_CHECKLIST_VAN_HANH.md", "09_DU_PHONG.md"]),
    ("PHẦN VI. NGUỒN, SỐ LIỆU VÀ GIỚI HẠN", ["11_NGUON_VA_GIOI_HAN.md"]),
]


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run(run, size=12, bold=None, italic=None, mono=False):
    # Academic profile requires Times New Roman throughout, including code.
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text, size=12):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run(paragraph.add_run(text[pos:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("`"):
            set_run(paragraph.add_run(token[1:-1]), size=max(9, size - 1), mono=True)
        elif token.startswith("**"):
            set_run(paragraph.add_run(token[2:-2]), size=size, bold=True)
        else:
            set_run(paragraph.add_run(token[1:-1]), size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        set_run(paragraph.add_run(text[pos:]), size=size)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run(run, size=10)


def configure(doc):
    setup_document(doc, profile="academic")
    section = doc.sections[0]
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 14, 14, 6),
        ("Heading 3", 12, 10, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.color.rgb = BLACK
        style.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("AI EXAM CONTROL · GIÁO TRÌNH 5 BUỔI")
    set_run(r, size=9, italic=True)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GIÁO ÁN VÀ SỔ TAY HỌC")
    set_run(r, size=18, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI THỊ GIÁC MÁY TÍNH QUA DỰ ÁN\nAI EXAM CONTROL")
    set_run(r, size=22, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tài liệu đào tạo 5 buổi · 90 phút mỗi buổi")
    set_run(r, size=14, italic=True)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [
        ("Đối tượng", "Học viên bắt đầu từ nền tảng phổ thông"),
        ("Trọng tâm", "Dữ liệu · Machine Learning · Deep Learning · Computer Vision"),
        ("Thực hành", "Gán nhãn · IoU · theo dõi thời gian · fine-tune · đánh giá"),
        ("Phiên bản", f"{date.today().strftime('%m/%Y')}"),
    ]
    for idx, (label, value) in enumerate(labels):
        table.cell(idx, 0).width = Cm(3.2)
        table.cell(idx, 1).width = Cm(10.5)
        for cell in table.rows[idx].cells:
            set_cell_margins(cell, 100, 100, 100, 100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(table.cell(idx, 0))
        p0 = table.cell(idx, 0).paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(p0.add_run(label), size=11, bold=True)
        p1 = table.cell(idx, 1).paragraphs[0]
        set_run(p1.add_run(value), size=11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Biên soạn từ mã nguồn, tài liệu và artifact thực tế của dự án"), size=11, italic=True)
    doc.add_page_break()


def add_front_matter(doc):
    doc.add_heading("HƯỚNG DẪN SỬ DỤNG", level=1)
    p = doc.add_paragraph()
    add_inline(p, "Tài liệu này xây kiến thức chuyên ngành từ chính bài toán của AI Exam Control. Học viên quan sát dữ liệu và hiện tượng trước, sau đó mới gọi tên khái niệm, thực hành và giải thích lại bằng lời của mình.")
    for text in (
        "Người dạy dùng Phần II làm kịch bản 90 phút cho từng buổi.",
        "Học viên dùng Phần III để ôn khái niệm và Phần IV để ghi sản phẩm thực hành.",
        "Phần V chỉ dùng sau khi học viên đã hiểu pipeline; không học thuộc câu trả lời mẫu.",
        "Các con số mô phỏng trong hoạt động dạy học không phải kết quả đo của hệ thống.",
        "Không sửa model, trọng số, cơ sở dữ liệu vận hành hoặc bằng chứng gốc trong quá trình học.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, text)

    doc.add_heading("MỤC LỤC NỘI DUNG", level=1)
    for title, files in SECTIONS:
        p = doc.add_paragraph()
        add_inline(p, title, 12)
        p.runs[0].bold = True
        for filename in files:
            first = first_heading(OUT_DIR / filename)
            p = doc.add_paragraph(style="List Bullet 2")
            add_inline(p, first)
    doc.add_page_break()


def first_heading(path):
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        if line.startswith("# "):
            return line[2:].strip()
    return path.stem


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        set_run(p.add_run(line), size=9, mono=True)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def add_table_from_rows(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    normalized = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    font_size = 9 if cols >= 5 else 10
    for ri, row in enumerate(normalized):
        prevent_row_split(table.rows[ri])
        if ri == 0:
            set_repeat_table_header(table.rows[ri])
        for ci, value in enumerate(row):
            cell = table.cell(ri, ci)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                shade_cell(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, font_size)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator_row(line):
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def add_markdown(doc, path):
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            rows = [parse_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table_from_rows(doc, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(3, len(heading.group(1)))
            p = doc.add_heading(level=level)
            add_inline(p, heading.group(2), 16 if level == 1 else 14 if level == 2 else 12)
            i += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", raw)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", raw)
        checkbox = re.match(r"^\s*[-*]\s+\[([ xX])\]\s+(.+)$", raw)
        if checkbox:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, ("☒ " if checkbox.group(1).lower() == "x" else "☐ ") + checkbox.group(2))
        elif bullet:
            indent = len(raw) - len(raw.lstrip())
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_inline(p, bullet.group(1))
        elif numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.right_indent = Cm(0.5)
            add_inline(p, stripped.lstrip("> "), 11)
            for run in p.runs:
                run.italic = True
        else:
            para_lines = [stripped]
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if not nxt or nxt.startswith(("#", "```", "|", ">")) or re.match(r"^\s*[-*+]\s+", lines[i]) or re.match(r"^\s*\d+[.)]\s+", lines[i]):
                    break
                para_lines.append(nxt)
                i += 1
            p = doc.add_paragraph()
            add_inline(p, " ".join(para_lines))
            continue
        i += 1
    if code_lines:
        add_code_block(doc, code_lines)


def main():
    doc = Document()
    configure(doc)
    add_cover(doc)
    add_front_matter(doc)
    for section_index, (section_title, files) in enumerate(SECTIONS):
        if section_index:
            doc.add_page_break()
        doc.add_heading(section_title, level=1)
        for filename in files:
            add_markdown(doc, OUT_DIR / filename)

    core = doc.core_properties
    core.title = "Giáo án và sổ tay học AI thị giác máy tính qua dự án AI Exam Control"
    core.subject = "Giáo trình 5 buổi về dữ liệu, ML, DL, CV, gán nhãn, fine-tune và đánh giá"
    core.author = "AI Exam Control"
    core.keywords = "AI, ML, Deep Learning, Computer Vision, YOLO, fine-tune, giáo án"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
